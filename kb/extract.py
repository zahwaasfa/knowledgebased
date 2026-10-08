"""Ekstraksi berkas fisik -> blok {h: heading, p: paragraf}. PDF, DOCX, PPTX, XLSX, CSV, TXT/MD."""
from __future__ import annotations
import csv
import re
from collections import Counter
from pathlib import Path
from .config import S

HEAD = [
    (1, re.compile(r"^(BAB|BAGIAN|CHAPTER)\s+[IVXLC\d]+\b.{0,100}$", re.I)),
    (4, re.compile(r"^\d{1,2}\.\d{1,2}\.\d{1,2}\.\d{1,2}[.)]?\s+\S.{1,100}$")),
    (3, re.compile(r"^\d{1,2}\.\d{1,2}\.\d{1,2}[.)]?\s+\S.{1,100}$")),
    (2, re.compile(r"^\d{1,2}\.\d{1,2}[.)]?\s+[A-Za-z]\S.{1,100}$")),
    (1, re.compile(r"^\d{1,2}[.)]\s+[A-Z][^.;:]{2,80}$")),
    (2, re.compile(r"^[A-Z][.)]\s+[A-Z][^.;]{2,80}$")),
]
TOC = re.compile(r"\.{4,}\s*\d+\s*$")


def heading_level(line: str) -> int:
    if len(line) > 110 or TOC.search(line) or line[-1] in ",;":
        return 0
    m = re.match(r"^(#{1,5})\s+\S", line)
    if m:
        return len(m.group(1))
    for lvl, rx in HEAD:
        if rx.match(line):
            return lvl
    letters = re.sub(r"[^A-Za-z]", "", line)
    if 4 <= len(line) <= 90 and len(letters) >= 4 and line.upper() == line and not line.endswith("."):
        return 1
    return 0


def _h(level, text, page=None): return {"kind": "h", "level": level, "text": text.lstrip("# ").strip(), "page": page}
def _p(text, page=None): return {"kind": "p", "text": text.strip(), "page": page, "atomic": False, "section": None, "item": None}


def text_to_blocks(text: str, page: int | None = None) -> list[dict]:
    lines = [l.strip() for l in text.splitlines()]
    wrap = max([len(l) for l in lines] + [1])
    blocks, buf = [], []

    def flush():
        if buf:
            blocks.append(_p(" ".join(buf), page)); buf.clear()
    for line in lines:
        if not line:
            flush(); continue
        if TOC.search(line):
            continue
        lvl = heading_level(line)
        if lvl:
            flush(); blocks.append(_h(lvl, line, page)); continue
        if buf and (re.match(r"^([-•●▪*]|\d+[.)]|[a-z][.)])\s", line) or (buf[-1][-1] in ".!?" and len(buf[-1]) < 0.75 * wrap and line[:1].isupper())):
            flush()
        buf.append(line)
    flush()
    return blocks


def _pdf(path: Path) -> list[dict]:
    from pypdf import PdfReader
    pages = [(p.extract_text() or "") for p in PdfReader(str(path)).pages]
    cnt = Counter()
    for t in pages:
        ls = [l.strip() for l in t.splitlines() if l.strip()]
        for l in set(ls[:2] + ls[-2:]):
            cnt[re.sub(r"\d+", "#", l)] += 1
    thr = max(3, int(0.4 * len(pages))) if len(pages) >= 4 else 10**9
    out = []
    for i, t in enumerate(pages, 1):
        ls = [l for l in t.splitlines() if not re.fullmatch(r"(halaman\s*)?\d{1,4}", l.strip(), re.I) and cnt[re.sub(r"\d+", "#", l.strip())] < thr]
        out += text_to_blocks("\n".join(ls), i)
    return out


def _rows(rows: list[list]) -> list[dict]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        return []
    head = [str(c).strip() for c in rows[0]]
    out = []
    for r in rows[1:S.max_table_rows + 1]:
        t = " | ".join(f"{h or f'kolom{i+1}'}: {str(v).strip()}" for i, (h, v) in enumerate(zip(head, r)) if str(v).strip())
        if t:
            b = _p(t); b["atomic"] = True; out.append(b)  # baris tabel = unit atomik, tidak dipotong
    return out


def _docx(path: Path) -> list[dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d, out = docx.Document(str(path)), []
    for ch in d.element.body.iterchildren():
        if ch.tag.endswith("}p"):
            p = Paragraph(ch, d); t = p.text.strip()
            if not t:
                continue
            m = re.match(r"(?:Heading|Judul)\s*(\d)", p.style.name or "", re.I)
            lvl = int(m.group(1)) if m else (heading_level(t) if len(t) < 100 else 0)
            out.append(_h(lvl, t) if lvl else _p(t))
        elif ch.tag.endswith("}tbl"):
            out += _rows([[c.text.strip() for c in r.cells] for r in Table(ch, d).rows])
    return out


def _pptx(path: Path) -> list[dict]:
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(str(path)).slides, 1):
        title = s.shapes.title.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        out.append(_h(1, f"Slide {i}: {title}".strip(": "), i))
        for sh in s.shapes:
            if sh.has_text_frame and sh != s.shapes.title:
                for para in sh.text_frame.paragraphs:
                    if para.text.strip():
                        out.append(_p(para.text, i))
    return out


def _xlsx(path: Path) -> list[dict]:
    import openpyxl
    wb, out = openpyxl.load_workbook(str(path), read_only=True, data_only=True), []
    for ws in wb.worksheets:
        out.append(_h(1, f"Lembar: {ws.title}"))
        out += _rows([["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True, max_row=S.max_table_rows + 1)])
    return out


def _csv(path: Path) -> list[dict]:
    raw = path.read_bytes()
    try:
        txt = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        txt = raw.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(txt[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [_h(1, f"Data: {path.stem}")] + _rows(list(csv.reader(txt.splitlines(), dialect)))


def extract_blocks(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    fn = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".csv": _csv}.get(ext)
    if fn:
        return fn(path)
    if ext in {".txt", ".md"}:
        return text_to_blocks(path.read_text(encoding="utf-8", errors="ignore"))
    raise ValueError(f"format tidak didukung: {ext}")


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def find_file(docs_dir: Path, file_name: str) -> Path | None:
    """Cocokkan nama berkas metadata dengan berkas di data/documents (abaikan ekstensi/tanda baca)."""
    if not file_name or not docs_dir.exists():
        return None
    stem = re.sub(r"\.[A-Za-z0-9]{2,4}$", "", file_name)
    want = _norm(stem)
    for f in sorted(docs_dir.rglob("*")):
        if f.is_file() and not f.name.startswith("."):
            have = _norm(f.stem)
            if have == want or (len(want) > 12 and (have.startswith(want) or want.startswith(have))):
                return f
    return None
